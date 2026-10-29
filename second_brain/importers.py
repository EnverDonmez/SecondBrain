import hashlib
import io
import json
import mimetypes
import stat
import zipfile
from pathlib import PurePosixPath

from .models import EventInput, MemoryInput, MemoryType, Principal, timestamp
from .security import Denied, require, source_identity

MAX_FILE = 100 * 1024 * 1024
MAX_EXPANDED = 300 * 1024 * 1024
TEXT_TYPES = {
    ".txt",
    ".md",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".html",
    ".css",
    ".sql",
    ".toml",
    ".yaml",
    ".yml",
    ".csv",
    ".json",
    ".jsonl",
    ".rs",
    ".go",
    ".java",
    ".c",
    ".cpp",
    ".xml",
    ".log",
}


def safe_zip(content):
    archive = zipfile.ZipFile(io.BytesIO(content))
    entries = archive.infolist()
    if len(entries) > 2000 or sum(e.file_size for e in entries) > MAX_EXPANDED:
        raise Denied("Archive exceeds extraction budget")
    for entry in entries:
        path = PurePosixPath(entry.filename.replace("\\", "/"))
        if (
            path.is_absolute()
            or ".." in path.parts
            or any(":" in p for p in path.parts)
            or stat.S_ISLNK(entry.external_attr >> 16)
        ):
            raise Denied("Unsafe archive path or symbolic link")
        if entry.file_size > MAX_FILE or entry.file_size / max(entry.compress_size, 1) > 1000:
            raise Denied("Archive entry exceeds size/ratio limits")
    return archive


class Importer:
    def __init__(self, db, events, memory, queue):
        self.db, self.events, self.memory, self.queue = db, events, memory, queue

    def submit(self, content, name, principal, scope="global", sensitivity=1):
        if len(content) > MAX_FILE or not content:
            raise Denied("File must be between 1 byte and 100 MiB")
        name = PurePosixPath(name.replace("\\", "/")).name[:250]
        ext = PurePosixPath(name).suffix.lower()
        if ext == ".zip":
            safe_zip(content).close()
        self.preflight_erasure(content, name, principal)
        if ext not in TEXT_TYPES | {
            ".zip",
            ".pdf",
            ".docx",
            ".png",
            ".jpg",
            ".jpeg",
            ".webp",
            ".wav",
            ".mp3",
            ".m4a",
            ".mp4",
            ".webm",
            ".ogg",
        }:
            raise Denied("File type is not supported")
        asset = self.events.store_asset(
            content,
            name,
            mimetypes.guess_type(name)[0] or "application/octet-stream",
            scope,
            sensitivity,
            principal,
        )
        return self.queue.enqueue("import", asset["id"], {"asset_id": asset["id"]}, principal)

    def preflight_erasure(self, content, name, principal):
        documents = []
        ext = PurePosixPath(name).suffix.lower()
        if ext == ".json":
            documents.append(json.loads(content))
        elif ext == ".zip":
            with safe_zip(content) as archive:
                for entry in archive.infolist():
                    if PurePosixPath(entry.filename).name == "conversations.json":
                        documents.append(json.loads(archive.read(entry)))
        erased = {
            r["target"]
            for r in self.db.rows(
                "SELECT target FROM tombstones WHERE tenant=? AND kind='source_key'", (principal.tenant,)
            )
        }
        for data in documents:
            if not isinstance(data, list):
                continue
            for conversation in data:
                if not isinstance(conversation, dict) or "mapping" not in conversation:
                    continue
                cid = str(
                    conversation.get("id")
                    or conversation.get("conversation_id")
                    or hashlib.sha256(json.dumps(conversation, sort_keys=True).encode()).hexdigest()
                )
                if any(
                    source_identity("chatgpt-export", cid + ":" + str(node_id)) in erased
                    for node_id in conversation["mapping"]
                ):
                    raise Denied(
                        "Export contains previously forgotten source messages. Import refused before storing original bytes."
                    )

    def process(self, job):
        principal = Principal(**job["payload"]["principal"])
        aid = job["payload"]["asset_id"]
        asset = self.db.one("SELECT * FROM assets WHERE id=? AND tenant=?", (aid, principal.tenant))
        if not asset:
            raise Denied("Import source was forgotten")
        require(principal, "capture", asset["scope"], asset["sensitivity"])
        content = self.events.vault.get(aid)
        result = self._parse(
            content,
            asset["name"],
            asset,
            principal,
            lambda: self.queue.heartbeat(job["id"], {"stage": "import"}),
        )
        with self.db.transaction() as con:
            con.execute(
                "INSERT OR REPLACE INTO checkpoints VALUES(?,?,?,?,?)",
                (principal.tenant, "import:" + aid, "complete", "1.0", timestamp(None)),
            )
        return result

    def _parse(self, content, name, asset, principal, progress=None):
        if progress:
            progress()
        ext = PurePosixPath(name).suffix.lower()
        if ext == ".zip":
            count, duplicates, unavailable, event_ids = 0, 0, [], []
            with safe_zip(content) as archive:
                for entry in archive.infolist():
                    if progress:
                        progress()
                    if entry.is_dir():
                        continue
                    basename = PurePosixPath(entry.filename).name
                    member = archive.read(entry)
                    if basename == "conversations.json":
                        imported = self.chatgpt(json.loads(member), asset, principal, progress)
                    elif PurePosixPath(basename).suffix.lower() in TEXT_TYPES | {".pdf", ".docx"}:
                        imported = self._parse(member, entry.filename, asset, principal, progress)
                    else:
                        # Entire original ZIP remains encrypted; do not claim transcription/OCR.
                        unavailable.append(entry.filename)
                        continue
                    count += imported.get("events", 0)
                    duplicates += imported.get("duplicates", 0)
                    event_ids += imported.get("event_ids", [])
            return {
                "events": count,
                "duplicates": duplicates,
                "event_ids": event_ids,
                "preserved_without_extraction": unavailable,
            }
        if ext == ".json":
            data = json.loads(content)
            if isinstance(data, list) and all(isinstance(c, dict) and "mapping" in c for c in data):
                return self.chatgpt(data, asset, principal, progress)
        if ext == ".jsonl":
            lines = [json.loads(line) for line in content.decode("utf-8-sig").splitlines() if line.strip()]
            if lines and all(isinstance(e, dict) and "source" in e and "source_id" in e for e in lines):
                results = []
                for e in lines:
                    if progress:
                        progress()
                    results.append(
                        self.events.ingest(
                            EventInput(
                                **{
                                    **e,
                                    "scope": asset["scope"],
                                    "sensitivity": asset["sensitivity"],
                                    "attachments": [asset["id"]],
                                }
                            ),
                            principal,
                        )
                    )
                return self._counts(results)
        pages = []
        if ext in TEXT_TYPES:
            pages = [("file", content.decode("utf-8-sig", errors="replace"))]
        elif ext == ".pdf":
            from pypdf import PdfReader

            pdf = PdfReader(io.BytesIO(content))
            if len(pdf.pages) > 2000:
                raise Denied("PDF exceeds page limit")
            pages = [(f"page:{i + 1}", p.extract_text() or "") for i, p in enumerate(pdf.pages)]
        elif ext == ".docx":
            from docx import Document

            safe_zip(content).close()
            document = Document(io.BytesIO(content))
            pages = [
                (
                    "document",
                    "\n".join(
                        [p.text for p in document.paragraphs]
                        + [" | ".join(c.text for c in r.cells) for t in document.tables for r in t.rows]
                    ),
                )
            ]
        elif ext in {".png", ".jpg", ".jpeg", ".webp"}:
            from PIL import Image

            with Image.open(io.BytesIO(content)) as image:
                metadata = {"width": image.width, "height": image.height, "format": image.format}
            event = self.events.ingest(
                EventInput(
                    source="file",
                    source_id=asset["hash"] + ":" + name,
                    actor="document",
                    event_type="file.image",
                    scope=asset["scope"],
                    sensitivity=asset["sensitivity"],
                    payload=metadata,
                    attachments=[asset["id"]],
                ),
                principal,
            )
            return {
                **self._counts([event]),
                "extraction": "metadata_only",
                "next_step": "Choose local OCR processing or attach a manual caption; original retained",
            }
        else:
            event = self.events.ingest(
                EventInput(
                    source="file",
                    source_id=asset["hash"] + ":" + name,
                    actor="document",
                    event_type="file.media",
                    scope=asset["scope"],
                    sensitivity=asset["sensitivity"],
                    payload={"name": name},
                    attachments=[asset["id"]],
                ),
                principal,
            )
            return {
                **self._counts([event]),
                "extraction": "original_preserved",
                "next_step": "Choose local media processing or attach a manual timestamped transcript",
            }
        results = []
        for location, text in pages:
            for i in range(0, max(len(text), 1), 8000):
                if progress:
                    progress()
                chunk = text[i : i + 8000]
                identity = hashlib.sha256(content).hexdigest() + ":" + name + ":" + location + ":" + str(i)
                result = self.events.ingest(
                    EventInput(
                        source="file",
                        source_id=identity,
                        actor="document",
                        event_type="file.extracted",
                        text=chunk,
                        scope=asset["scope"],
                        sensitivity=asset["sensitivity"],
                        metadata={"name": name, "location": location, "offset": i},
                        attachments=[asset["id"]],
                    ),
                    principal,
                )
                results.append(result)
                if result["id"] and chunk.strip():
                    self.memory.propose(
                        MemoryInput(
                            memory_type=MemoryType.EXTERNAL_KNOWLEDGE,
                            subject=name,
                            predicate=location + ":" + str(i),
                            value=chunk,
                            scope=asset["scope"],
                            sensitivity=asset["sensitivity"],
                            source_event_ids=[result["id"]],
                            confidence=1,
                        ),
                        principal,
                    )
        return self._counts(results)

    @staticmethod
    def _counts(results):
        return {
            "events": sum(not r["duplicate"] for r in results),
            "duplicates": sum(r["duplicate"] for r in results),
            "event_ids": [r["id"] for r in results if r["id"]],
        }

    def chatgpt(self, data, asset, principal, progress=None):
        if not isinstance(data, list):
            raise ValueError("Expected ChatGPT conversations array")
        results = []
        for conversation in data:
            cid = str(
                conversation.get("id")
                or conversation.get("conversation_id")
                or hashlib.sha256(json.dumps(conversation, sort_keys=True).encode()).hexdigest()
            )
            mapping = conversation.get("mapping", {})
            for node_id, node in mapping.items():
                if progress:
                    progress()
                message = node.get("message")
                if not message:
                    continue
                role = message.get("author", {}).get("role", "system")
                parts = message.get("content", {}).get("parts", [])
                text = "\n".join(p for p in parts if isinstance(p, str))
                event = EventInput(
                    source="chatgpt-export",
                    source_id=cid + ":" + str(node_id),
                    actor=role if role in {"user", "assistant", "tool", "system"} else "system",
                    event_type="conversation.message",
                    text=text,
                    occurred_at=timestamp(message["create_time"])
                    if message.get("create_time") is not None
                    else None,
                    conversation_id=cid,
                    session_id=cid,
                    parent_id=str(node["parent"]) if node.get("parent") else None,
                    scope=asset["scope"],
                    sensitivity=asset["sensitivity"],
                    metadata={
                        "title": conversation.get("title"),
                        "children": node.get("children", []),
                        "current_node": conversation.get("current_node"),
                        "conversation_metadata": {k: v for k, v in conversation.items() if k != "mapping"},
                    },
                    payload=node,
                    attachments=[asset["id"]],
                )
                result = self.events.ingest(event, principal)
                results.append(result)
                if (
                    result["id"]
                    and not result["duplicate"]
                    and event.actor == "user"
                    and {"read", "propose"} <= set(principal.permissions)
                ):
                    self.extract_explicit(result["id"], principal)
        return self._counts(results)

    def extract_explicit(self, event_id, principal):
        from .preferences import preference_candidates

        event = self.events.get(event_id, principal)
        if event["actor"] != "user":
            return {
                "candidates": [],
                "reason": "Only direct user assertions can propose preferences or rules",
            }
        candidates = [
            self.memory.propose(
                MemoryInput(
                    **item,
                    subject="user",
                    source_event_ids=[event_id],
                    scope=event["scope"],
                    sensitivity=event["sensitivity"],
                ),
                principal,
            )
            for item in preference_candidates(event["text"])
        ]
        return {"candidates": candidates, "approval_required": True}
